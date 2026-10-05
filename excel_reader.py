import openpyxl
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import CellCoordinatesException


def format_row_ranges(rows):
    if not rows:
        return "Không có"

    ranges = []
    range_start = rows[0]
    range_end = rows[0]

    for row in rows[1:]:
        if row == range_end + 1:
            range_end = row
            continue

        ranges.append(str(range_start) if range_start == range_end else f"{range_start}-{range_end}")
        range_start = range_end = row

    ranges.append(str(range_start) if range_start == range_end else f"{range_start}-{range_end}")
    return ", ".join(ranges)


def summarize_data_column(sheet, column="B", first_row=4):
    data_rows = [
        row
        for row in range(first_row, sheet.max_row + 1)
        if sheet[f"{column}{row}"].value is not None
        and str(sheet[f"{column}{row}"].value).strip()
    ]

    if not data_rows:
        return f"Không tìm thấy dữ liệu ở cột {column} từ dòng {first_row} trở xuống."

    return (
        f"Các dòng sẽ xử lý ở cột {column}: {format_row_ranges(data_rows)}.\n"
        f"Tổng số mã hàng: {len(data_rows)}."
    )


def parse_start_cell(cell_address):
    normalized_address = cell_address.strip().replace("$", "").upper()
    if not normalized_address:
        raise ValueError("Vui lòng nhập ô bắt đầu, ví dụ B4.")

    try:
        code_column, start_row = coordinate_from_string(normalized_address)
        code_column_index = column_index_from_string(code_column)
    except (ValueError, CellCoordinatesException) as exc:
        raise ValueError("Ô bắt đầu không hợp lệ. Vui lòng nhập theo dạng B4, B18, ...") from exc

    if start_row < 1:
        raise ValueError("Dòng bắt đầu phải lớn hơn hoặc bằng 1.")

    if code_column_index > 16382:
        raise ValueError("Cột mã hàng phải đứng trước cột XFC để xác định được cột số lượng.")

    quantity_column = get_column_letter(code_column_index + 2)
    canonical_address = f"{code_column}{start_row}"
    return canonical_address, code_column, quantity_column, start_row


def to_quantity(value):
    """Chuyển giá trị ô Excel thành số; ô trống là 0, không phải số thì trả về None."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).replace(",", "").strip()
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return None


def collect_excel_data(sheet, code_column, quantity_column, start_row):
    data_list = []

    for current_row in range(start_row, sheet.max_row + 1):
        code = sheet[f"{code_column}{current_row}"].value

        # Chỉ xử lý ô mã hàng có dữ liệu, tự bỏ qua mọi dòng trống xen giữa
        if code is None or not str(code).strip():
            continue

        data_list.append({
            "code": str(code).strip(),
            "qty": to_quantity(sheet[f"{quantity_column}{current_row}"].value),
            "row": current_row,
        })

    return data_list


def open_sheet(file_path, sheet_name=""):
    """Mở workbook và trả về sheet cần dùng; sheet trống nghĩa là sheet đang active."""
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    if not sheet_name:
        return workbook.active
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Không tìm thấy Sheet '{sheet_name}'.")
    return workbook[sheet_name]


def find_duplicates(items):
    """Các mã xuất hiện nhiều hơn một lần (không phân biệt hoa/thường): [(mã, [dòng, ...])]."""
    rows_by_code = {}
    display_code = {}
    for item in items:
        key = item["code"].lower()
        display_code.setdefault(key, item["code"])
        rows_by_code.setdefault(key, []).append(item["row"])

    return [
        (display_code[key], rows)
        for key, rows in rows_by_code.items()
        if len(rows) > 1
    ]


def precheck(items):
    """Kiểm tra dữ liệu Excel trước khi chạy, không đụng đến web."""
    return {
        "total": len(items),
        "duplicates": find_duplicates(items),
        "invalid": [(i["code"], i["row"]) for i in items if i["qty"] is None],
        "zero": [i["row"] for i in items if i["qty"] == 0],
    }


def describe_precheck(report):
    """Biến kết quả precheck thành danh sách (mức độ, nội dung) để hiển thị."""
    lines = [("info", f"Tổng số mã hàng: {report['total']}.")]

    for code, rows in report["duplicates"]:
        lines.append(("warn", f"Mã trùng: {code} ở các dòng {format_row_ranges(rows)} (số lượng sẽ cộng dồn lên web)."))
    for code, row in report["invalid"]:
        lines.append(("error", f"Số lượng không phải số: mã {code} (dòng {row}) sẽ bị bỏ qua."))
    if report["zero"]:
        lines.append(("warn", f"Số lượng bằng 0 ở các dòng: {format_row_ranges(report['zero'])}."))

    if len(lines) == 1:
        lines.append(("ok", "Không phát hiện vấn đề nào trong dữ liệu Excel."))
    return lines
