import pytest
from openpyxl import Workbook

import excel_reader as er


def test_parse_start_cell():
    assert er.parse_start_cell("b4") == ("B4", "B", "D", 4)
    assert er.parse_start_cell("$C$10") == ("C10", "C", "E", 10)
    with pytest.raises(ValueError):
        er.parse_start_cell("")
    with pytest.raises(ValueError):
        er.parse_start_cell("4B")


def test_format_row_ranges():
    assert er.format_row_ranges([]) == "Không có"
    assert er.format_row_ranges([4, 5, 6, 9, 11, 12]) == "4-6, 9, 11-12"


def test_to_quantity():
    assert er.to_quantity(None) == 0
    assert er.to_quantity("") == 0
    assert er.to_quantity(3) == 3.0
    assert er.to_quantity("1,200") == 1200.0
    assert er.to_quantity("-2.5") == -2.5
    assert er.to_quantity("abc") is None


def make_sheet():
    ws = Workbook().active
    ws["B4"], ws["D4"] = "A1", 2
    ws["B6"], ws["D6"] = " B2 ", -1
    ws["B7"] = "C3"
    ws["B8"], ws["D8"] = "a1", "x"
    return ws


def test_collect_excel_data_skips_blank_rows_and_flags_bad_quantity():
    assert er.collect_excel_data(make_sheet(), "B", "D", 4) == [
        {"code": "A1", "qty": 2.0, "row": 4},
        {"code": "B2", "qty": -1.0, "row": 6},
        {"code": "C3", "qty": 0.0, "row": 7},
        {"code": "a1", "qty": None, "row": 8},
    ]


def test_summarize_data_column():
    text = er.summarize_data_column(make_sheet(), "B", 4)
    assert "4, 6-8" in text and "Tổng số mã hàng: 4" in text
    assert "Không tìm thấy" in er.summarize_data_column(make_sheet(), "B", 20)


def test_find_duplicates_is_case_insensitive():
    items = er.collect_excel_data(make_sheet(), "B", "D", 4)
    assert er.find_duplicates(items) == [("A1", [4, 8])]


def test_precheck_and_description():
    report = er.precheck(er.collect_excel_data(make_sheet(), "B", "D", 4))
    assert report["total"] == 4
    assert report["invalid"] == [("a1", 8)]
    assert report["zero"] == [7]
    levels = [level for level, _ in er.describe_precheck(report)]
    assert "warn" in levels and "error" in levels


def test_describe_precheck_clean_file():
    lines = er.describe_precheck({"total": 2, "duplicates": [], "invalid": [], "zero": []})
    assert lines[-1][0] == "ok"


def test_open_sheet(tmp_path):
    path = tmp_path / "t.xlsx"
    wb = Workbook()
    wb.active.title = "S1"
    wb.save(path)
    assert er.open_sheet(str(path), "S1").title == "S1"
    assert er.open_sheet(str(path), "").title == "S1"
    with pytest.raises(ValueError):
        er.open_sheet(str(path), "Khong co")
