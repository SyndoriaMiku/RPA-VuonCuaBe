import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import tkinter.scrolledtext as scrolledtext
import openpyxl
import time
import threading
import os
import subprocess
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import CellCoordinatesException
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys


def format_row_ranges(rows):
    if not rows:
        return "Không có"

    ranges = []
    range_start = rows[0]
    range_end = rows[0]

    for row in rows[1:]:
        if row == range_end + 1:
            range_end = row
            continue

        ranges.append(str(range_start) if range_start == range_end else f"{range_start}-{range_end}")
        range_start = range_end = row

    ranges.append(str(range_start) if range_start == range_end else f"{range_start}-{range_end}")
    return ", ".join(ranges)


def summarize_data_column(sheet, column="B", first_row=4):
    data_rows = [
        row
        for row in range(first_row, sheet.max_row + 1)
        if sheet[f"{column}{row}"].value is not None
        and str(sheet[f"{column}{row}"].value).strip()
    ]

    if not data_rows:
        return f"Không tìm thấy dữ liệu ở cột {column} từ dòng {first_row} trở xuống."

    return (
        f"Các dòng sẽ xử lý ở cột {column}: {format_row_ranges(data_rows)}.\n"
        f"Tổng số mã hàng: {len(data_rows)}."
    )


def collect_excel_data(sheet, code_column, quantity_column, start_row):
    data_list = []

    for current_row in range(start_row, sheet.max_row + 1):
        code = sheet[f"{code_column}{current_row}"].value

        # Chỉ xử lý ô mã hàng có dữ liệu, tự bỏ qua mọi dòng trống xen giữa
        if code is None or not str(code).strip():
            continue

        qty_val = sheet[f"{quantity_column}{current_row}"].value
        qty = float(qty_val) if qty_val is not None else 0

        data_list.append({
            "code": str(code).strip(),
            "qty": qty,
            "row": current_row,
        })

    return data_list


def choose_file():
    file_path = filedialog.askopenfilename(
        title="Chọn File Excel",
        filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")]
    )
    if file_path:
        entry_file_path.delete(0, tk.END)
        entry_file_path.insert(0, file_path)
        
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet_name = entry_sheet.get()
            if sheet_name and sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
            else:
                sheet = wb.active

            c2_value = sheet["C2"].value
            if c2_value:
                lbl_title.config(text=str(c2_value))
                root.title(f"Auto Robot - {c2_value}")
            else:
                lbl_title.config(text="[Ô C2 Trống]")
                root.title("Auto Robot")

            try:
                _, scan_column, _, scan_start_row = parse_start_cell(entry_start_cell.get())
            except ValueError:
                scan_column, scan_start_row = "B", 4

            data_summary = summarize_data_column(sheet, scan_column, scan_start_row)
            log_summary = (
                f"\n--- ĐÃ TẢI FILE THÀNH CÔNG ---\n"
                f"File: {os.path.basename(file_path)}\n"
                f"Sheet: {sheet.title}\n"
                f"{data_summary}\n"
            )
            append_log(log_summary)
        except Exception:
            lbl_title.config(text="[Lỗi đọc ô C2]")
            root.title("Auto Robot")
            messagebox.showerror("Lỗi đọc Excel", "Không thể đọc thông tin từ file Excel đã chọn.")

def parse_start_cell(cell_address):
    normalized_address = cell_address.strip().replace("$", "").upper()
    if not normalized_address:
        raise ValueError("Vui lòng nhập ô bắt đầu, ví dụ B4.")

    try:
        code_column, start_row = coordinate_from_string(normalized_address)
        code_column_index = column_index_from_string(code_column)
    except (ValueError, CellCoordinatesException) as exc:
        raise ValueError("Ô bắt đầu không hợp lệ. Vui lòng nhập theo dạng B4, B18, ...") from exc

    if start_row < 1:
        raise ValueError("Dòng bắt đầu phải lớn hơn hoặc bằng 1.")

    if code_column_index > 16382:
        raise ValueError("Cột mã hàng phải đứng trước cột XFC để xác định được cột số lượng.")

    quantity_column = get_column_letter(code_column_index + 2)
    canonical_address = f"{code_column}{start_row}"
    return canonical_address, code_column, quantity_column, start_row


def start_automation_thread():
    try:
        start_cell_config = parse_start_cell(entry_start_cell.get())
    except ValueError as exc:
        messagebox.showerror("Lỗi", str(exc))
        return

    # Làm mờ nút chạy để ngăn người dùng bấm nhiều lần
    btn_run.config(state=tk.DISABLED)
    
    # Tạo một luồng phụ để chạy tool
    thread = threading.Thread(target=run_automation_worker, args=(start_cell_config,))
    thread.daemon = True # Đảm bảo luồng phụ tự tắt khi bạn ấn X đóng phần mềm
    thread.start()

def open_chrome_debug_window():
    chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
    if not os.path.exists(chrome_path):
        messagebox.showerror("Lỗi", f"Không tìm thấy Chrome tại:\n{chrome_path}")
        return
    try:
        os.system("taskkill /f /im chromedriver.exe >nul 2>&1")
    except:
        pass

    try:
        subprocess.Popen(
            [
                chrome_path,
                "--remote-debugging-port=9999",
                "--user-data-dir=C:\\ChromeDebug",
            ]
        )
        append_log("Đã mở Chrome Debug (port 9999).\n")
    except Exception as e:
        messagebox.showerror("Lỗi", f"Không thể mở Chrome Debug:\n{e}")

def run_automation_worker(start_cell_config):
    try:
        run_automation(start_cell_config)
    finally:
        # Cập nhật giao diện từ main thread của Tkinter
        root.after(0, lambda: btn_run.config(state=tk.NORMAL))

def append_log(message):
    txt_log.config(state=tk.NORMAL)
    txt_log.insert(tk.END, message)
    txt_log.see(tk.END)
    txt_log.config(state=tk.DISABLED)

def copy_selected_log(event=None):
    try:
        selected_text = txt_log.selection_get()
    except tk.TclError:
        return "break"

    root.clipboard_clear()
    root.clipboard_append(selected_text)
    return "break"

def select_all_log(event=None):
    txt_log.tag_add(tk.SEL, "1.0", tk.END)
    txt_log.mark_set(tk.INSERT, "1.0")
    txt_log.see(tk.INSERT)
    return "break"

def run_automation(start_cell_config):
    file_path = entry_file_path.get()
    sheet_name = entry_sheet.get()
    start_cell, code_column, quantity_column, start_row = start_cell_config
    try:
        sleep_time = float(combo_sleep.get())
    except ValueError:
        sleep_time = 2.5
    
    if not file_path:
        messagebox.showerror("Lỗi", "Vui lòng chọn file Excel trước khi chạy.")
        return
    
    # ---------------------------------------------------------
    # PHẦN 1: ĐỌC TOÀN BỘ DỮ LIỆU EXCEL TỪ Ô ĐƯỢC CHỌN
    # ---------------------------------------------------------
    data_list = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if sheet_name:
            if sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
            else:
                messagebox.showerror("Lỗi", f"Không tìm thấy Sheet '{sheet_name}'.")
                return
        else:
            sheet = wb.active
            
        data_list = collect_excel_data(sheet, code_column, quantity_column, start_row)
            
    except Exception as e:
        messagebox.showerror("Lỗi Đọc Excel", f"Có lỗi xảy ra khi đọc file:\n{e}")
        return

    if not data_list:
        messagebox.showwarning("Cảnh báo", f"Không tìm thấy dữ liệu nào bắt đầu từ ô {start_cell}.")
        return

    # ---------------------------------------------------------
    # PHẦN 2: CHẠY VÒNG LẶP TRÊN WEB (SELENIUM)
    # ---------------------------------------------------------
    log_messages = []
    
    try:
        chrome_options = Options()
        chrome_options.add_experimental_option("debuggerAddress", "127.0.0.1:9999")
        driver = webdriver.Chrome(options=chrome_options)
        
        # --- BẢN VÁ MỚI: TỰ ĐỘNG CHUYỂN ĐÚNG TAB KIOTVIET ---
        for handle in driver.window_handles:
            driver.switch_to.window(handle)
            if "kiotviet.vn" in driver.current_url.lower():
                break
        # ----------------------------------------------------

        wait = WebDriverWait(driver, 10)
        
        append_log(f"\n----------------------------------------\n")
        append_log(f"Đọc Excel từ {start_cell}; số lượng tại cột {quantity_column}.\n")
        append_log(f"Bắt đầu chạy {len(data_list)} mã hàng...\n")

        # Vòng lặp duyệt qua từng dòng dữ liệu lấy từ Excel
        for item in data_list:
            code = item["code"]
            qty_excel = item["qty"]
            
            try:
                # Step 1: Tìm ô input và click vào nó trước
                search_input = wait.until(EC.element_to_be_clickable((By.ID, "productSearchInput")))
                
                # Giả lập thao tác click của người dùng
                search_input.click()
                time.sleep(0.5)
                
                search_input.send_keys(Keys.CONTROL + "a")
                search_input.send_keys(Keys.DELETE)
                search_input.send_keys(code)
                
                # Chờ web load kết quả mới dựa theo thiết lập UI
                time.sleep(sleep_time) 
                
                # Step 2: Click vào kết quả bằng Javascript (Né AdGuard)
                first_result = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, ".output-complete ul li")))
                driver.execute_script("arguments[0].click();", first_result)
                
                # Step 3: Đợi input số lượng hiện ra
                input_td = wait.until(EC.presence_of_element_located((
                    By.CSS_SELECTOR, "table[role='grid'] tbody tr:nth-child(1) td.cell-order-number input"
                )))
                
                # Step 4: Lấy số lượng trên web
                list_cells = driver.find_elements(By.CSS_SELECTOR, "table[role='grid'] tbody tr:nth-child(1) td.cell-quantity.txtR")
                txt_web_value = list_cells[0].text.strip()
                txt_web_value = txt_web_value.replace(",", "").replace(" ", "")
                qty_web = float(txt_web_value) if txt_web_value else 0
                
                # Step 5: Tính tổng và kiểm tra điều kiện
                total_value = qty_web + qty_excel
                
                if total_value < 0:
                    error_msg = f"- Mã hàng {code} số lượng cân chỉnh không hợp lệ (Tổng: {total_value})"
                    log_messages.append(error_msg)
                    append_log(error_msg + "\n")
                    continue 
                
                # Step 6: Điền số lượng bằng API AngularJS
                driver.execute_script("""
                    var elm = arguments[0];
                    var val = arguments[1];
                    var $el = angular.element(elm);
                    $el.val(val).triggerHandler('input');
                    $el.triggerHandler('change');
                    $el.triggerHandler('blur');
                """, input_td, total_value)
                
                
                time.sleep(1) 
                
            except Exception as e_row:
                
                chi_tiet_loi = str(e_row)
                if "Stacktrace" in chi_tiet_loi or "Timeout" in chi_tiet_loi:
                    error_msg = f"- BỎ QUA: Mã {code} - Không tìm thấy sản phẩm trên KiotViet."
                else:
                    error_msg = f"- LỖI Web: Mã {code}. Chi tiết: {chi_tiet_loi.split(chr(10))[0]}"
                
                log_messages.append(error_msg)
                append_log(error_msg + "\n")
                continue

        # Sau khi chạy xong toàn bộ danh sách
        append_log("--- HOÀN TẤT QUÁ TRÌNH ---\n")

    except Exception as e:
        messagebox.showerror("Lỗi Hệ Thống", f"Không thể kết nối với Web:\n{e}")

# --- THIẾT KẾ GIAO DIỆN (GUI) ---
root = tk.Tk()
root.title("Auto Robot")
root.geometry("850x650")
root.eval('tk::PlaceWindow . center')

lbl_title = tk.Label(root, text="Tiêu đề C2 sẽ hiển thị ở đây", font=("Arial", 14, "bold"), fg="#D32F2F")
lbl_title.grid(row=0, column=0, columnspan=3, pady=15)

tk.Label(root, text="File Excel:").grid(row=1, column=0, padx=10, pady=10, sticky="e")
entry_file_path = tk.Entry(root, width=45)
entry_file_path.grid(row=1, column=1, padx=5, pady=10)
btn_browse = tk.Button(root, text="Chọn File", command=choose_file)
btn_browse.grid(row=1, column=2, padx=5, pady=10)

tk.Label(root, text="Tên Sheet (bỏ trống=mặc định):").grid(row=2, column=0, padx=10, pady=10, sticky="e")
entry_sheet = tk.Entry(root, width=20)
entry_sheet.insert(0, "Cân Chỉnh Hàng Hóa")
entry_sheet.grid(row=2, column=1, padx=5, pady=10, sticky="w")

frame_sleep = tk.Frame(root)
frame_sleep.grid(row=2, column=2, padx=5, pady=10, sticky="w")
tk.Label(frame_sleep, text="Chờ load web (giây):").pack(side=tk.LEFT)
combo_sleep = ttk.Combobox(frame_sleep, values=["1.5", "2.0", "2.5", "3.0", "3.5", "4.0", "4.5", "5.0"], width=5)
combo_sleep.set("2.5")
combo_sleep.pack(side=tk.LEFT)

tk.Label(root, text="Ô bắt đầu mã hàng:").grid(row=3, column=0, padx=10, pady=10, sticky="e")
entry_start_cell = tk.Entry(root, width=10)
entry_start_cell.insert(0, "B4")
entry_start_cell.grid(row=3, column=1, padx=5, pady=10, sticky="w")

btn_run = tk.Button(root, text="▶ CHẠY TỰ ĐỘNG", command=start_automation_thread, bg="#4CAF50", fg="white", font=("Arial", 11, "bold"))
btn_run.grid(row=4, column=0, columnspan=2, pady=15, ipadx=20, ipady=5)

btn_open_chrome = tk.Button(root, text="Mở Chrome Debug", command=open_chrome_debug_window, font=("Arial", 10, "bold"))
btn_open_chrome.grid(row=4, column=2, padx=5, pady=15, ipadx=10, ipady=5)

# Khung Log
tk.Label(root, text="Nhật ký chạy (Log):").grid(row=5, column=0, padx=10, sticky="nw")
txt_log = scrolledtext.ScrolledText(root, width=65, height=10, fg="red", bg="#f9f9f9")
txt_log.grid(row=6, column=0, columnspan=3, padx=10, pady=5)
txt_log.config(state=tk.DISABLED, insertwidth=0, insertontime=0, insertofftime=0, takefocus=1)
txt_log.bind("<Control-c>", copy_selected_log)
txt_log.bind("<Control-a>", select_all_log)
txt_log.bind("<Button-1>", lambda event: txt_log.focus_set())

root.mainloop()
